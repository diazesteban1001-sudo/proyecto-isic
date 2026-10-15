#!/usr/bin/env python3
"""
test_md_a_docx_comentarios.py — control de los comentarios HTML (<!-- … -->)
en md_a_docx.py.

El defecto
----------
md_a_docx.py no quitaba los comentarios HTML. El borrador anterior no tenía
ninguno, así que no se notaba; el vigente lleva 203, los <!-- F: … --> de
trazabilidad y las notas de sección, y salían como texto visible en 122
párrafos del .docx (2026-10-15).

Borrador sintético
------------------
  A. Un comentario de una línea, en su propia línea, entre dos frases de un
     mismo párrafo: tiene que desaparecer, y las dos frases tienen que quedar
     en un solo párrafo.
  B. Un comentario de varias líneas: tiene que desaparecer entero.
  C. Un comentario en medio de una línea: tiene que desaparecer, y el texto de
     los dos lados quedarse.
  D. Un «<!--» dentro de un bloque de código cercado: tiene que quedarse.
  E. En todo el .docx, el único «<!--» es el del bloque de código, y no queda
     ningún «F:».
  F. Un comentario con texto detrás en su misma línea, y otro en la línea
     siguiente: el texto entre los dos se queda.

Dos controles positivos, por el corolario de la regla 6 de CLAUDE.md: un
chequeo que no puede fallar no vale nada.
  - Con el filtro desactivado —quitar_comentarios_html reemplazado por la
    identidad—, tienen que fallar todos los casos de comentario, A, B, C, E y
    F. D no depende del filtro.
  - Con un mutante cuyo patrón del comentario en sus propias líneas sí cruza
    un «-->», tiene que fallar F: se traga el texto entre los dos.

Uso:
    .venv/bin/python .claude/skills/sintesis-consultoria/scripts/test_md_a_docx_comentarios.py
Devuelve 0 si los casos se comportan como se espera.
"""

import os
import re
import sys
import tempfile

from docx import Document

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
import md_a_docx  # noqa: E402

BORRADOR = """# Informe sintético

Primera frase del párrafo.
<!-- F: MARCA-UNA, outputs/x.json > campo -->
Segunda frase del mismo párrafo.

Antes del comentario <!-- F: MARCA-EN-MEDIO --> y después.

<!-- MARCA-VARIAS-1
MARCA-VARIAS-2
MARCA-VARIAS-3, F: con su fuente -->

Texto después del comentario de varias líneas.

```
codigo con <!-- esto se queda -->
```

<!-- F: MARCA-F1 --> Texto entre comentarios.
<!-- F: MARCA-F2 -->
"""


def texto_del_docx(ruta):
    doc = Document(ruta)
    parrafos = [p.text for p in doc.paragraphs]
    for tabla in doc.tables:
        for fila in tabla.rows:
            parrafos.extend(c.text for c in fila.cells)
    return parrafos


def casos(parrafos):
    todo = "\n".join(parrafos)
    return [
        ("A. comentario de una línea, fuera",
         "MARCA-UNA" not in todo),
        ("A. las dos frases, en un solo párrafo",
         "Primera frase del párrafo. Segunda frase del mismo párrafo." in parrafos),
        ("B. comentario de varias líneas, fuera",
         "MARCA-VARIAS" not in todo),
        ("C. comentario en medio de una línea, fuera, y el texto se queda",
         "MARCA-EN-MEDIO" not in todo
         and any(p.startswith("Antes del comentario") and p.endswith("y después.") for p in parrafos)),
        ("D. «<!--» dentro del bloque de código, se queda",
         "<!-- esto se queda -->" in todo),
        ("E. un solo «<!--» en el .docx, y ningún «F:»",
         todo.count("<!--") == 1 and "F:" not in todo),
        ("F. el texto entre dos comentarios se queda",
         "Texto entre comentarios." in todo and "MARCA-F" not in todo),
    ]


# El patrón sin la guarda: con DOTALL, «.*?» se extiende hasta el primer «-->»
# seguido de fin de línea, aunque haya otro «-->» antes.
SIN_GUARDA = re.compile(r"^[ \t]*<!--.*?-->[ \t]*(?:\n|\Z)", re.DOTALL | re.MULTILINE)


def correr(filtro=True, patron=None):
    original = md_a_docx.quitar_comentarios_html
    patron_original = md_a_docx.COMENTARIO_EN_SUS_LINEAS
    if not filtro:
        md_a_docx.quitar_comentarios_html = lambda texto: texto
    if patron is not None:
        md_a_docx.COMENTARIO_EN_SUS_LINEAS = patron
    try:
        with tempfile.TemporaryDirectory() as d:
            entrada = os.path.join(d, "borrador.md")
            salida = os.path.join(d, "borrador.docx")
            with open(entrada, "w", encoding="utf-8") as f:
                f.write(BORRADOR)
            md_a_docx.convertir(entrada, salida)
            return casos(texto_del_docx(salida))
    finally:
        md_a_docx.quitar_comentarios_html = original
        md_a_docx.COMENTARIO_EN_SUS_LINEAS = patron_original


def main():
    con = correr(filtro=True)
    print("Con el filtro:")
    for titulo, ok in con:
        print(f"  [{'OK' if ok else 'FALLA'}] {titulo}")

    sin = correr(filtro=False)
    print("Control positivo, sin el filtro (cada caso de comentario tiene que fallar):")
    for titulo, ok in sin:
        print(f"  [{'pasa' if ok else 'falla'}] {titulo}")

    mut = correr(patron=SIN_GUARDA)
    f_mutante = dict(mut)["F. el texto entre dos comentarios se queda"]
    print("Control positivo, mutante sin la guarda:")
    print(f"  [{'pasa' if f_mutante else 'falla'}] F. el texto entre dos comentarios se queda")

    todos_bien = all(ok for _, ok in con)
    # D no depende del filtro: sin él, el «<!--» del bloque también se queda.
    control_falla = all(not ok for titulo, ok in sin if not titulo.startswith("D."))
    print(f"\nCon el filtro, todos los casos: {'sí' if todos_bien else 'NO'}")
    print(f"Sin el filtro, fallan todos los casos de comentario: {'sí' if control_falla else 'NO'}")
    print(f"Sin la guarda, falla F: {'sí' if not f_mutante else 'NO'}")
    return 0 if todos_bien and control_falla and not f_mutante else 1


if __name__ == "__main__":
    sys.exit(main())
